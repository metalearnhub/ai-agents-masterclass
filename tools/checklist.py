"""Usage:
  python3 tools/checklist.py next                     -> next post that needs NO screen recording (JSON), or {"none":true}
  python3 tools/checklist.py publish "<tab>" <#> <url> <YYYY-MM-DD>
Queue order: recording 'No' before 'Optional'; Website Blog Plan before Botcake Series; then Priority order."""
import openpyxl,sys,json,datetime,os
F=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),'checklist','Blog content checklist.xlsx')
wb=openpyxl.load_workbook(F)
def rows(ws):
    hdr={ws.cell(4,c).value:c for c in range(1,ws.max_column+1)}
    for r in range(5,ws.max_row+1):
        if ws.cell(r,1).value is None: continue
        yield r,hdr,{k:ws.cell(r,c).value for k,c in hdr.items()}
if sys.argv[1]=='next':
    cand=[]
    for ti,tab in enumerate(['Website Blog Plan','Botcake Series']):
        for r,hdr,v in rows(wb[tab]):
            rec=str(v.get('Screen recording needed from Dijae') or '')
            if v.get('Status')!='Planned': continue
            rank=0 if rec.startswith('No') else 1 if rec.startswith('Optional') else None
            if rank is None: continue
            cand.append((rank,ti,v['Priority order'],tab,v))
    cand.sort(key=lambda x:x[:3])
    print(json.dumps({'tab':cand[0][3],**cand[0][4]},default=str,indent=1) if cand else '{"none":true}')
elif sys.argv[1]=='publish':
    tab,num,url,date=sys.argv[2],int(sys.argv[3]),sys.argv[4],sys.argv[5]
    ws=wb[tab]
    for r,hdr,v in rows(ws):
        if v['#']==num:
            ws.cell(r,hdr['Status']).value='Published'
            c=ws.cell(r,hdr['Blog URL']); c.value=url; c.hyperlink=url
            d=ws.cell(r,hdr['Publish date']); d.value=datetime.date.fromisoformat(date); d.number_format='yyyy-mm-dd'
            wb.save(F); print('updated',tab,num); break
    else: sys.exit('row not found')
